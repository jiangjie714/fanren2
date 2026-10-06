#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
软著申请材料的 Word（.docx）版生成器，与 tools/gen-softcopyright.mjs 的 PDF 版本内容完全一致。

用法（需 python>=3.9 并已安装 python-docx，推荐在虚拟环境中运行）：
    python3 tools/gen-softcopyright-docx.py <source|manual|all> --owner "XX 科技有限公司"

口径：
  源程序文档 —— 每页 50 行（按「ASCII 0.6em / CJK 1.0em」折算折行占位分页），提交前 30 页 + 后 30 页；
  用户手册   —— 17 章 ×（文字页 + 截图页），图文并茂。
  页眉：软件全称 + 版本号（右上页码）；页脚：著作权人全称。
"""
import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

from PIL import Image
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MJS = os.path.join(ROOT, 'tools', 'gen-softcopyright.mjs')

LINES_PER_PAGE = 50
UNITS_PER_LINE = 60  # ≈100 个 ASCII 字符


def line_units(line):
    u = 0.0
    for ch in line:
        u += 1.0 if ('⺀' <= ch <= '鿿' or '＀' <= ch <= '￯' or '　' <= ch <= '〿') else 0.6
    return u


# --------------------------------------------------------------------------- 基础工具

def set_run_font(run, name, size, east_asia=None):
    run.font.name = name
    run.font.size = Pt(size)
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    if east_asia:
        rFonts.set(qn('w:eastAsia'), east_asia)


def add_field(paragraph, instr, size=9, name='宋体'):
    """插入 Word 域（如 PAGE），用于自动页码。"""
    run = paragraph.add_run()
    set_run_font(run, name, size)
    f1 = OxmlElement('w:fldChar')
    f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = instr
    f2 = OxmlElement('w:fldChar')
    f2.set(qn('w:fldCharType'), 'end')
    run._r.append(f1)
    run._r.append(it)
    run._r.append(f2)


def no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'none')
        el.set(qn('w:sz'), '0')
        borders.append(el)
    tblPr.append(borders)


def setup_page(doc, header_text, footer_text, content_width_pt, first_page_plain=False):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(1.7)
    sec.right_margin = Cm(1.7)
    sec.header_distance = Cm(0.6)
    sec.footer_distance = Cm(0.6)

    if first_page_plain:
        sec.different_first_page_header_footer = True

    def fill(header):
        p = header.paragraphs[0]
        p.paragraph_format.tab_stops.add_tab_stop(Pt(content_width_pt), WD_TAB_ALIGNMENT.RIGHT)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header_text)
        set_run_font(r, '宋体', 9)
        r2 = p.add_run('\t第 ')
        set_run_font(r2, '宋体', 9)
        add_field(p, 'PAGE')
        r3 = p.add_run(' 页')
        set_run_font(r3, '宋体', 9)

    if first_page_plain:
        fp = sec.first_page_header
        fp.paragraphs[0].text = ''
        ff = sec.first_page_footer
        ff.paragraphs[0].text = ''
        fill(sec.header)
    else:
        fill(sec.header)

    fp = sec.footer.paragraphs[0]
    fp.paragraph_format.space_after = Pt(0)
    r = fp.add_run(footer_text)
    set_run_font(r, '宋体', 9)
    return sec


def page_break(paragraph):
    paragraph.add_run().add_break(WD_BREAK.PAGE)


# --------------------------------------------------------------------------- 源程序文档

def source_order():
    src = open(MJS, encoding='utf-8').read()
    m = re.search(r'const SOURCE_ORDER = \[(.*?)\];', src, re.S)
    return re.findall(r"'([^']+\.ts)'", m.group(1))


def collect_lines():
    lines = []
    for rel in source_order():
        path = os.path.join(ROOT, rel)
        if not os.path.exists(path):
            print('[warn] 缺失', rel)
            continue
        lines.append('// ==================== 源文件：%s ====================' % rel)
        with open(path, encoding='utf-8') as f:
            for raw in f:
                line = raw.rstrip()
                t = line.strip()
                if not t or t.startswith('//') or t.startswith('/*') or t.startswith('*'):
                    continue
                lines.append(line.rstrip())
    return lines


def paginate(lines):
    pages, cur, cost = [], [], 0.0
    for line in lines:
        c = max(1, int(-(-line_units(line) // UNITS_PER_LINE)))
        if cost + c > LINES_PER_PAGE and cur:
            pages.append(cur)
            cur, cost = [], 0.0
        cur.append(line)
        cost += c
    if cur:
        pages.append(cur)
    return pages


def gen_source(args):
    lines = collect_lines()
    pages = paginate(lines)
    head = pages[:30]
    tail = pages[-30:] if len(pages) > 60 else []
    picked = head + tail

    doc = Document()
    content_width = (21.0 - 1.7 * 2) * 28.3465  # cm → pt
    setup_page(doc, '%s %s' % (args.name, args.version), '著作权人：%s' % args.owner, content_width)

    styles = doc.styles['Normal']
    styles.font.name = 'Courier New'
    styles.font.size = Pt(8)

    for pi, page in enumerate(picked):
        for li, line in enumerate(page):
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(12.5)
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            r = p.add_run(line)
            set_run_font(r, 'Courier New', 8, east_asia='宋体')
            if li == len(page) - 1 and pi != len(picked) - 1:
                p.add_run().add_break(WD_BREAK.PAGE)

    out = os.path.join(args.outdir, '源程序_%s_%s.docx' % (args.name, args.version))
    doc.save(out)
    print('[source] 源文件 %d 个，有效代码行 %d，全部 %d 页；提交前 %d + 后 %d = %d 页'
          % (len(source_order()), len(lines), len(pages), len(head), len(tail), len(picked)))
    print('[source] 输出', out)


# --------------------------------------------------------------------------- 用户手册

def load_manual():
    js = ("import { MANUAL } from '%s'; console.log(JSON.stringify(MANUAL));"
          % os.path.join(ROOT, 'tools', 'softcopyright-manual.mjs'))
    res = subprocess.run(['node', '--input-type=module', '-e', js], capture_output=True, text=True, cwd=ROOT)
    if res.returncode != 0:
        print(res.stderr, file=sys.stderr)
        raise SystemExit('读取手册内容失败')
    return json.loads(res.stdout)


def shrink(img_path, cache_dir, target_width=760):
    """把截图按显示尺寸重采样并转 JPEG，控制 docx 体积（打印清晰度仍足够）。"""
    name = os.path.splitext(os.path.basename(img_path))[0] + '.jpg'
    out = os.path.join(cache_dir, name)
    if os.path.exists(out):
        return out
    im = Image.open(img_path).convert('RGB')
    if im.width > target_width:
        im = im.resize((target_width, int(im.height * target_width / im.width)), Image.LANCZOS)
    im.save(out, 'JPEG', quality=88, optimize=True)
    return out


def gen_manual(args):
    manual = load_manual()
    cache_dir = tempfile.mkdtemp(prefix='rz-img-')
    doc = Document()
    content_width = (21.0 - 1.7 * 2) * 28.3465
    setup_page(doc, '%s %s' % (args.name, args.version), '著作权人：%s' % args.owner,
               content_width, first_page_plain=True)

    styles = doc.styles['Normal']
    styles.font.name = '宋体'
    styles.font.size = Pt(11)

    def heading(text, size):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(text)
        r.bold = True
        set_run_font(r, '宋体', size, east_asia='宋体')
        return p

    # 封面
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(args.name)
    r.bold = True
    set_run_font(r, '宋体', 28, east_asia='宋体')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    r = p.add_run('用 户 手 册')
    set_run_font(r, '宋体', 20, east_asia='宋体')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(60)
    r = p.add_run('版本号：%s' % args.version)
    set_run_font(r, '宋体', 13, east_asia='宋体')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('著作权人：%s' % args.owner)
    set_run_font(r, '宋体', 13, east_asia='宋体')
    page_break(doc.add_paragraph())

    # 目录
    heading('目 录', 15)
    for c in manual:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(c['h'])
        set_run_font(r, '宋体', 12, east_asia='宋体')
    p = doc.add_paragraph()
    r = p.add_run('附：版本信息')
    set_run_font(r, '宋体', 12, east_asia='宋体')
    page_break(doc.add_paragraph())

    # 正文
    for ci, c in enumerate(manual):
        heading(c['h'], 15)
        for text in c['p']:
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            pf.line_spacing = 1.3
            pf.space_after = Pt(4)
            if re.match(r'^\d+\.\d+\s', text):
                pf.first_line_indent = Pt(0)
            else:
                pf.first_line_indent = Pt(22)
            r = p.add_run(text)
            if re.match(r'^\d+\.\d+\s', text):
                r.bold = True
            set_run_font(r, '宋体', 11, east_asia='宋体')
        page_break(doc.add_paragraph())

        imgs = [(f, cap) for f, cap in c['imgs'] if os.path.exists(os.path.join(ROOT, 'docs', 'screenshots', f))]
        if not imgs:
            continue
        heading('%s　界面示意' % c['h'], 15)
        cols = 1 if len(imgs) == 1 else 2
        rows = (len(imgs) + cols - 1) // cols
        table = doc.add_table(rows=rows, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        no_borders(table)
        for gi, (f, cap) in enumerate(imgs):
            cell = table.cell(gi // cols, gi % cols)
            cell.paragraphs[0].text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            width = Cm(11.0) if cols == 1 else Cm(5.2)
            p.add_run().add_picture(shrink(os.path.join(ROOT, 'docs', 'screenshots', f), cache_dir), width=width)
            cap_p = cell.add_paragraph()
            cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_p.paragraph_format.space_after = Pt(6)
            r = cap_p.add_run('图 %d-%d　%s' % (ci + 1, gi + 1, cap))
            set_run_font(r, '宋体', 9, east_asia='宋体')
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run('注：以上界面截图均为本软件实际运行画面，具体数值以软件内实际显示为准。')
        set_run_font(r, '宋体', 9, east_asia='宋体')
        page_break(doc.add_paragraph())

    # 版本信息
    heading('附：版本信息', 15)
    for text in ['软件全称：%s' % args.name,
                 '版 本 号：%s' % args.version,
                 '著作权人：%s' % args.owner,
                 '编程语言：TypeScript 5.4',
                 '源程序量：10673 行（64 个 TypeScript 源文件）',
                 '运行环境：抖音小游戏；Android 8.0 及以上 / iOS 12.0 及以上',
                 '开发完成日期：见《计算机软件著作权登记申请表》',
                 '本手册所述界面与操作以软件实际运行效果为准。']:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        set_run_font(r, '宋体', 11, east_asia='宋体')

    out = os.path.join(args.outdir, '用户手册_%s_%s.docx' % (args.name, args.version))
    doc.save(out)
    shutil.rmtree(cache_dir, ignore_errors=True)
    print('[manual] 章节 %d 个，输出 %s' % (len(manual), out))


# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('cmd', nargs='?', default='all', choices=['source', 'manual', 'all'])
    ap.add_argument('--name', default='凡人开仙缘游戏软件')
    ap.add_argument('--version', default='V1.0')
    ap.add_argument('--owner', default='【请填写：企业全称】')
    ap.add_argument('--outdir', default=os.path.join(ROOT, 'docs', '软著申请材料'))
    args = ap.parse_args()
    os.makedirs(args.outdir, exist_ok=True)
    if args.cmd in ('source', 'all'):
        gen_source(args)
    if args.cmd in ('manual', 'all'):
        gen_manual(args)


if __name__ == '__main__':
    main()
